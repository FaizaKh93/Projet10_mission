# tests/test_load_excel_to_db.py
"""Test fonctionnel de l'ingestion (scripts/load_excel_to_db.py).

Gratuit : on fabrique un petit fichier Excel dans un dossier temporaire, on lance
la vraie chaîne Excel -> validation Pydantic -> base SQLite, puis on relit la base.
Aucun appel réseau, aucune dépendance au fichier NBA réel.

Le fichier fabriqué **reproduit la corruption** du fichier réel : l'en-tête de la
colonne 3PM y est un `datetime.time(15, 0)`, comme l'a écrit Excel.
"""
import datetime
import sqlite3
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

# scripts/ n'est pas un package installé (comme src/ et eval/) - cf. conftest.py
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

from db_models import Stat
from load_excel_to_db import COLONNES, lire_donnees_nba, main

# Les deux joueurs servent aussi à vérifier les valeurs après insertion
JOUEURS = [
    {"Player": "Shai Gilgeous-Alexander", "Team": "OKC", "Age": 26, "PTS": 2485, "3PM": 160, "3PA": 433},
    {"Player": "Nikola Jokić", "Team": "DEN", "Age": 29, "PTS": 2072, "3PM": 140, "3PA": 329},
]


def construire_excel(chemin: Path, joueurs=JOUEURS, corrompre: bool = True) -> Path:
    """Fabrique un Excel minimal au format du fichier réel.

    Deux particularités reproduites fidèlement :
    - les données commencent à la 2e ligne (une ligne de titre au-dessus, d'où
      `header=1` à la lecture) ;
    - l'en-tête de 3PM est une **vraie cellule de type heure**.

    Ce second point impose openpyxl : si on écrit `datetime.time(15, 0)` via
    pandas, la valeur est enregistrée comme la chaîne "15:00:00" et relue comme
    du texte — ce qui ne reproduit pas la corruption. Il faut une cellule dont le
    type Excel est réellement une heure, comme celle du fichier d'origine.
    """
    colonnes = {"Player": [j["Player"] for j in joueurs], "Team": [j["Team"] for j in joueurs]}
    for xl in COLONNES:  # toutes les colonnes attendues par le script, à 0 par défaut
        colonnes[xl] = [j.get(xl, 0) for j in joueurs]
    colonnes["3PM"] = [j["3PM"] for j in joueurs]

    df = pd.DataFrame(colonnes)
    equipes = pd.DataFrame({"Code": ["OKC", "DEN"], "Nom": ["Oklahoma City Thunder", "Denver Nuggets"]})

    with pd.ExcelWriter(chemin) as writer:
        # startrow=1 : laisse une ligne de titre au-dessus, comme le fichier réel
        df.to_excel(writer, sheet_name="Données NBA", index=False, startrow=1)
        equipes.to_excel(writer, sheet_name="Equipe", index=False)

    if corrompre:
        # Remplace l'en-tête "3PM" par une valeur horaire, comme l'a fait Excel
        classeur = load_workbook(chemin)
        feuille = classeur["Données NBA"]
        colonne_3pm = list(df.columns).index("3PM") + 1  # openpyxl indexe à partir de 1
        feuille.cell(row=2, column=colonne_3pm).value = datetime.time(15, 0)
        classeur.save(chemin)

    return chemin


@pytest.fixture
def base_chargee(tmp_path):
    """Lance la vraie ingestion sur un Excel fabriqué, renvoie le chemin de la base."""
    excel = construire_excel(tmp_path / "test.xlsx")
    db = tmp_path / "test.db"
    main(excel=excel, url=f"sqlite:///{db}", season="2024-25")
    return db


# --- Lecture de l'Excel ---


def test_lecture_renomme_l_entete_corrompu(tmp_path):
    """L'en-tête `datetime.time(15, 0)` doit redevenir '3PM' à la lecture.

    Sans ce traitement, la règle de renommage générale produirait « 15_00_00 » —
    un identifiant SQL invalide, puisqu'il commence par un chiffre.
    """
    df = lire_donnees_nba(construire_excel(tmp_path / "test.xlsx"))
    assert "3PM" in df.columns
    assert not any(isinstance(c, datetime.time) for c in df.columns)


def test_lecture_echoue_si_entete_non_corrompu(tmp_path):
    """Si un futur export Excel ne présente plus la corruption, le script doit
    s'arrêter : cela signifierait que la structure du fichier a changé, et donc
    que la correspondance des colonnes n'est plus garantie."""
    chemin = construire_excel(tmp_path / "sain.xlsx", corrompre=False)
    with pytest.raises(RuntimeError, match="en-tête corrompu"):
        lire_donnees_nba(chemin)


# --- Ingestion de bout en bout ---


def test_ingestion_compte_les_lignes(base_chargee):
    con = sqlite3.connect(base_chargee)
    assert con.execute("SELECT COUNT(*) FROM teams").fetchone()[0] == 2
    assert con.execute("SELECT COUNT(*) FROM players").fetchone()[0] == 2
    assert con.execute("SELECT COUNT(*) FROM stats").fetchone()[0] == 2


def test_ingestion_conserve_les_valeurs(base_chargee):
    """Les valeurs relues doivent être identiques à celles du fichier source,
    y compris celles de la colonne dont l'en-tête était corrompu."""
    con = sqlite3.connect(base_chargee)
    ligne = con.execute(
        "SELECT s.pts_total, s.three_pm_total, s.three_pa_total FROM stats s "
        "JOIN players p USING(player_id) WHERE p.full_name LIKE 'Shai%'"
    ).fetchone()
    assert ligne == (2485, 160, 433)


def test_ingestion_relie_joueur_et_equipe(base_chargee):
    """La jointure doit fonctionner : c'est elle qui permettra au tool SQL de
    répondre « Denver Nuggets » plutôt que « DEN »."""
    con = sqlite3.connect(base_chargee)
    nom = con.execute(
        "SELECT t.name FROM stats s JOIN teams t ON t.code = s.team_code "
        "JOIN players p USING(player_id) WHERE p.full_name LIKE 'Nikola%'"
    ).fetchone()[0]
    assert nom == "Denver Nuggets"


def test_ingestion_rejouable(tmp_path):
    """Relancer l'ingestion ne doit pas dupliquer les lignes : le script repart
    d'une base vide (drop_all), puisqu'elle est entièrement dérivée de l'Excel."""
    excel = construire_excel(tmp_path / "test.xlsx")
    db = tmp_path / "test.db"
    main(excel=excel, url=f"sqlite:///{db}", season="2024-25")
    main(excel=excel, url=f"sqlite:///{db}", season="2024-25")
    con = sqlite3.connect(db)
    assert con.execute("SELECT COUNT(*) FROM stats").fetchone()[0] == 2


# --- Cohérence entre le script et le modèle ---


def test_correspondance_colonnes_coherente_avec_le_modele():
    """Garde-fou contre une divergence silencieuse.

    Le dictionnaire COLONNES (script) et le modèle Stat (base) contiennent chacun
    une liste de ~41 noms de colonnes, tenues à jour à la main. Si l'une évolue
    sans l'autre, l'ingestion planterait au premier INSERT — ou pire, oublierait
    une statistique sans rien signaler. Ce test compare les deux listes.
    """
    colonnes_modele = {c.name for c in Stat.__table__.columns}
    colonnes_script = set(COLONNES.values())
    gerees_automatiquement = {"stat_id", "player_id", "team_code", "season"}

    manquantes = colonnes_modele - colonnes_script - gerees_automatiquement
    en_trop = colonnes_script - colonnes_modele

    assert not manquantes, f"colonnes du modèle jamais remplies par l'ingestion : {manquantes}"
    assert not en_trop, f"colonnes insérées mais absentes du modèle : {en_trop}"
